"""Bounded parsers that treat uploaded content as inert data."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import Path
from uuid import uuid4

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from .models import (
    NormalizedDocument, NormalizedRecord, NormalizedSegment, SourceLocation,
)


class DocumentParseError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


MEDIA_TYPES = {
    "csv": "text/csv", "txt": "text/plain", "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def _decode(data: bytes) -> str:
    if b"\x00" in data:
        raise DocumentParseError("INVALID_ENCODING", "text input contains binary NUL bytes")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentParseError("INVALID_ENCODING", "text input must be valid UTF-8") from exc


def _zip_kind(data: bytes, max_uncompressed_bytes: int) -> str | None:
    if not data.startswith(b"PK"):
        return None
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            total = sum(item.file_size for item in archive.infolist())
            if total > max_uncompressed_bytes:
                raise DocumentParseError("EXPANSION_LIMIT", "document expands beyond the parsing limit")
            if any(item.file_size > max(1, item.compress_size) * 200 for item in archive.infolist()):
                raise DocumentParseError("EXPANSION_LIMIT", "document has an unsafe compression ratio")
            if "xl/workbook.xml" in names:
                return "xlsx"
            if "word/document.xml" in names:
                return "docx"
    except DocumentParseError:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise DocumentParseError("MALFORMED_DOCUMENT", "invalid ZIP-based office document") from exc
    return None


def detect_type(filename: str, data: bytes, max_uncompressed_bytes: int) -> str:
    suffix = Path(filename).suffix.casefold()
    if data.startswith(b"%PDF-"):
        detected = "pdf"
    else:
        detected = _zip_kind(data, max_uncompressed_bytes)
        if detected is None:
            text = _decode(data)
            first = next((line for line in text.splitlines() if line.strip()), "")
            detected = "csv" if suffix == ".csv" or ("," in first and len(first.split(",")) > 1) else "txt"
    expected = {".csv": "csv", ".txt": "txt", ".pdf": "pdf", ".xlsx": "xlsx", ".docx": "docx"}.get(suffix)
    if expected is None:
        raise DocumentParseError("UNSUPPORTED_TYPE", "supported extensions are csv, txt, pdf, xlsx, and docx")
    if expected != detected:
        raise DocumentParseError("TYPE_MISMATCH", "file content does not match its extension")
    return detected


def _source(document_id: str, version: int, locator: str, text: str, source_group: str, **kwargs) -> SourceLocation:
    return SourceLocation(document_id=document_id, version=version, locator=locator,
                          source_text=text[:8_000], source_group=source_group, **kwargs)


def _parse_csv(document_id: str, tenant_id: str, version: int, data: bytes, source_group: str) -> NormalizedDocument:
    text = _decode(data)
    try:
        rows = list(csv.DictReader(io.StringIO(text, newline="")))
    except csv.Error as exc:
        raise DocumentParseError("MALFORMED_DOCUMENT", "CSV parsing failed") from exc
    if not rows or not rows[0]:
        raise DocumentParseError("EMPTY_CONTENT", "CSV has no data rows")
    records = []
    for index, row in enumerate(rows, start=2):
        fields = {str(k).strip(): str(v or "").strip() for k, v in row.items() if k is not None}
        rendered = ", ".join(f"{key}={value}" for key, value in fields.items())
        location = _source(document_id, version, f"row:{index}", rendered, source_group, row=index)
        records.append(NormalizedRecord(record_id=f"{document_id}:row:{index}", location=location, fields=fields))
    return NormalizedDocument(document_id=document_id, tenant_id=tenant_id, source_type="csv", records=records)


def _parse_txt(document_id: str, tenant_id: str, version: int, data: bytes, source_group: str) -> NormalizedDocument:
    text = _decode(data)
    segments = [NormalizedSegment(
        segment_id=f"{document_id}:line:{index}",
        location=_source(document_id, version, f"line:{index}", line, source_group, line=index), text=line)
        for index, line in enumerate(text.splitlines(), start=1) if line.strip()]
    if not segments:
        raise DocumentParseError("EMPTY_CONTENT", "text document has no readable content")
    return NormalizedDocument(document_id=document_id, tenant_id=tenant_id, source_type="txt", segments=segments)


def _parse_pdf(document_id: str, tenant_id: str, version: int, data: bytes, source_group: str) -> NormalizedDocument:
    try:
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise DocumentParseError("ENCRYPTED_DOCUMENT", "encrypted PDFs are not supported")
        segments = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                segments.append(NormalizedSegment(
                    segment_id=f"{document_id}:page:{page_number}",
                    location=_source(document_id, version, f"page:{page_number}", text, source_group,
                                     page=page_number), text=text[:100_000]))
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError("MALFORMED_DOCUMENT", "PDF parsing failed") from exc
    if not segments:
        raise DocumentParseError("OCR_REQUIRED", "PDF contains no extractable text; bounded OCR review is required")
    return NormalizedDocument(document_id=document_id, tenant_id=tenant_id, source_type="pdf", segments=segments)


def _parse_xlsx(document_id: str, tenant_id: str, version: int, data: bytes, source_group: str) -> NormalizedDocument:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_vba=False, keep_links=False)
        records = []
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
            if not any(headers):
                continue
            for row_number, values in enumerate(rows, start=2):
                fields = {headers[i]: str(value).strip() for i, value in enumerate(values)
                          if i < len(headers) and headers[i] and value is not None}
                if not fields:
                    continue
                rendered = ", ".join(f"{key}={value}" for key, value in fields.items())
                location = _source(document_id, version, f"sheet:{sheet.title}:row:{row_number}",
                                   rendered, source_group, sheet=sheet.title, row=row_number)
                records.append(NormalizedRecord(
                    record_id=f"{document_id}:{sheet.title}:{row_number}", location=location, fields=fields))
        workbook.close()
    except Exception as exc:
        raise DocumentParseError("MALFORMED_DOCUMENT", "XLSX parsing failed") from exc
    if not records:
        raise DocumentParseError("EMPTY_CONTENT", "spreadsheet has no readable records")
    return NormalizedDocument(document_id=document_id, tenant_id=tenant_id, source_type="xlsx", records=records)


def _parse_docx(document_id: str, tenant_id: str, version: int, data: bytes, source_group: str) -> NormalizedDocument:
    try:
        document = DocxDocument(io.BytesIO(data))
        segments = []
        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if text:
                segments.append(NormalizedSegment(
                    segment_id=f"{document_id}:paragraph:{index}",
                    location=_source(document_id, version, f"paragraph:{index}", text, source_group), text=text))
        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                text = " | ".join(cell.text.strip() for cell in row.cells)
                if text.strip(" |"):
                    segments.append(NormalizedSegment(
                        segment_id=f"{document_id}:table:{table_index}:row:{row_index}",
                        location=_source(document_id, version,
                                         f"table:{table_index}:row:{row_index}", text, source_group,
                                         row=row_index), text=text))
    except Exception as exc:
        raise DocumentParseError("MALFORMED_DOCUMENT", "DOCX parsing failed") from exc
    if not segments:
        raise DocumentParseError("EMPTY_CONTENT", "DOCX has no readable paragraphs or table rows")
    return NormalizedDocument(document_id=document_id, tenant_id=tenant_id, source_type="docx", segments=segments)


def parse_document(*, filename: str, data: bytes, tenant_id: str, document_id: str,
                   version: int, source_group: str, max_uncompressed_bytes: int) -> tuple[str, NormalizedDocument]:
    source_type = detect_type(filename, data, max_uncompressed_bytes)
    parser = {"csv": _parse_csv, "txt": _parse_txt, "pdf": _parse_pdf,
              "xlsx": _parse_xlsx, "docx": _parse_docx}[source_type]
    return source_type, parser(document_id, tenant_id, version, data, source_group)
