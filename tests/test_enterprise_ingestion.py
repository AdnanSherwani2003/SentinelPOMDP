from __future__ import annotations

import io
import secrets

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from pydantic import ValidationError
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from fastapi.testclient import TestClient

from sentinel_pomdp.domains import DomainAction, DomainPack, DomainPackService, DomainPackStatus
from sentinel_pomdp.graph.repository import InMemoryGraphRepository
from sentinel_pomdp.ingestion import (
    CandidateStatus, DatasetConfiguration, EnterpriseDataService, IngestionError,
    LocalDocumentStore, ReviewAction,
)
from sentinel_pomdp.ingestion.parsers import DocumentParseError, parse_document
from sentinel_pomdp.intelligence import DeterministicFakeLLM, InvestigationIntelligenceService
from sentinel_pomdp.api.app import build_dependencies, create_app
from sentinel_pomdp.api.settings import AppSettings
from sentinel_pomdp.security import Permission


def dataset() -> DatasetConfiguration:
    return DatasetConfiguration(
        dataset_id="delivery", name="Delivery records",
        field_mapping={"purchase_order": "po", "supplier": "supplier",
                       "expected_date": "expected", "actual_date": "actual"})


def pack(*, sources: int = 2) -> DomainPack:
    outcomes_delayed = {
        "EVIDENCE_SUPPORTED": .90, "EVIDENCE_CONFLICTED": .02,
        "EVIDENCE_CONTRADICTED": .02, "EVIDENCE_NOT_FOUND": .04,
        "DUPLICATE_INFORMATION": .01, "SOURCE_HASH_MISMATCH": .01}
    outcomes_ontime = {
        "EVIDENCE_SUPPORTED": .10, "EVIDENCE_CONFLICTED": .05,
        "EVIDENCE_CONTRADICTED": .60, "EVIDENCE_NOT_FOUND": .23,
        "DUPLICATE_INFORMATION": .01, "SOURCE_HASH_MISMATCH": .01}
    return DomainPack(
        domain_id="delivery_delay", name="Delivery delay", claim_type="shipment_timeliness",
        hypotheses=["delayed", "on_time"], prior={"delayed": .55, "on_time": .45},
        actions=[DomainAction(
            action_id="inspect_delivery", relation_hypotheses={"EXPERIENCED_DELAY": "delayed"},
            cost=.08)],
        observation_model={"inspect_delivery": {
            "delayed": outcomes_delayed, "on_time": outcomes_ontime}},
        answer_threshold=.70, minimum_independent_sources=sources,
        status=DomainPackStatus.APPROVED, model_version="reviewed-2026-10",
        calibration_note="Illustrative reviewer-authored values; not benchmark calibrated.",
        approved_by="reviewer-1")


@pytest.fixture
def services(tmp_path):
    repository = InMemoryGraphRepository()
    store = LocalDocumentStore(tmp_path / "company")
    ingestion = EnterpriseDataService(store, repository, max_upload_bytes=100_000)
    ingestion.configure_dataset("tenant-a", dataset())
    return store, repository, ingestion


def test_parsers_preserve_csv_xlsx_docx_and_text_locations(tmp_path):
    common = dict(tenant_id="tenant-a", document_id="doc-1", version=1,
                  source_group="source-a", max_uncompressed_bytes=1_000_000)
    _, csv_doc = parse_document(
        filename="a.csv", data=b"po,expected\nPO-1,2026-01-01\n", **common)
    assert csv_doc.records[0].location.row == 2
    assert csv_doc.records[0].location.source_text == "po=PO-1, expected=2026-01-01"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Receiving"
    sheet.append(["po", "actual"])
    sheet.append(["PO-1", "2026-01-03"])
    xlsx = io.BytesIO()
    workbook.save(xlsx)
    _, book_doc = parse_document(filename="a.xlsx", data=xlsx.getvalue(), **common)
    assert book_doc.records[0].location.sheet == "Receiving"
    assert book_doc.records[0].location.row == 2

    word = DocxDocument()
    word.add_paragraph("PO-1 expected 2026-01-01 and delivered 2026-01-03")
    docx = io.BytesIO()
    word.save(docx)
    _, word_doc = parse_document(filename="a.docx", data=docx.getvalue(), **common)
    assert word_doc.segments[0].location.locator == "paragraph:1"

    _, text_doc = parse_document(filename="a.txt", data=b"line one\nline two\n", **common)
    assert text_doc.segments[1].location.line == 2


def test_pdf_without_text_requires_explicit_ocr_review():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    content = io.BytesIO()
    writer.write(content)
    with pytest.raises(DocumentParseError, match="OCR") as caught:
        parse_document(filename="scan.pdf", data=content.getvalue(), tenant_id="tenant-a",
                       document_id="doc", version=1, source_group="scan",
                       max_uncompressed_bytes=1_000_000)
    assert caught.value.code == "OCR_REQUIRED"


def test_pdf_text_is_extracted_with_page_provenance():
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=100)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 50 Td (PO-7 expected 2026-01-01 delivered 2026-01-03) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    content = io.BytesIO()
    writer.write(content)
    _, parsed = parse_document(
        filename="delivery.pdf", data=content.getvalue(), tenant_id="tenant-a",
        document_id="doc", version=1, source_group="pdf-source",
        max_uncompressed_bytes=1_000_000)
    assert "PO-7" in parsed.segments[0].text
    assert parsed.segments[0].location.page == 1


def test_encrypted_pdf_is_rejected_without_attempting_decryption():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("not-a-runtime-secret")
    content = io.BytesIO()
    writer.write(content)
    with pytest.raises(DocumentParseError) as caught:
        parse_document(
            filename="encrypted.pdf", data=content.getvalue(), tenant_id="tenant-a",
            document_id="doc", version=1, source_group="pdf-source",
            max_uncompressed_bytes=1_000_000)
    assert caught.value.code == "ENCRYPTED_DOCUMENT"


def test_type_spoofing_and_bounded_upload_are_rejected(services):
    _, _, ingestion = services
    with pytest.raises(IngestionError) as mismatch:
        ingestion.upload("tenant-a", dataset_id="delivery", filename="fake.pdf",
                         content=b"po,actual\nPO-1,2026-01-03")
    assert mismatch.value.code == "TYPE_MISMATCH"
    with pytest.raises(IngestionError) as too_large:
        ingestion.upload("tenant-a", dataset_id="delivery", filename="large.txt",
                         content=b"x" * 100_001)
    assert too_large.value.code == "UPLOAD_TOO_LARGE"
    with pytest.raises(IngestionError) as bad_name:
        ingestion.upload("tenant-a", dataset_id="delivery", filename="bad\r\nname.csv",
                         content=b"po,actual\nPO-1,2026-01-03")
    assert bad_name.value.code == "INVALID_FILENAME"


def test_review_publication_lookup_tenant_isolation_and_withdrawal(services):
    store, repository, ingestion = services
    expected = ingestion.upload(
        "tenant-a", dataset_id="delivery", filename="expected.csv",
        content=b"po,supplier,expected,actual\nPO-100,Supplier A,2026-10-01,\n")
    actual = ingestion.upload(
        "tenant-a", dataset_id="delivery", filename="actual.csv",
        content=b"po,supplier,expected,actual\nPO-100,Supplier A,,2026-10-04\n")
    assert expected.candidates_created == 0
    assert actual.candidates_created == 1
    candidate = store.list_candidates("tenant-a")[0]
    assert candidate.status == CandidateStatus.NEEDS_REVIEW
    assert {source.location.document_id for source in candidate.sources} == {
        expected.document.document_id, actual.document.document_id}

    approved = ingestion.review(
        "tenant-a", candidate.candidate_id, ReviewAction.APPROVE, "reviewer-1")
    assert approved.status == CandidateStatus.APPROVED
    chain = repository.get_provenance_chain(approved.graph_fact_id, "tenant-a")
    assert chain is not None and len(chain.supporting_evidence) == 2
    assert len({item.source_group for item in chain.supporting_evidence}) == 2
    assert repository.get_provenance_chain(approved.graph_fact_id, "tenant-b") is None
    lookup = ingestion.lookup("tenant-a", "Was PO-100 delayed?")
    assert lookup["status"] == "FOUND" and len(lookup["records"]) == 1

    ingestion.withdraw("tenant-a", actual.document.document_id)
    assert ingestion.lookup("tenant-a", "Was PO-100 delayed?")["status"] == "NO_APPROVED_EVIDENCE"
    assert len(repository.get_fact_evidence(approved.graph_fact_id, "tenant-a")[0]) == 1


def test_duplicate_and_version_supersession(services):
    store, _, ingestion = services
    first_bytes = b"po,supplier,expected,actual\nPO-2,S,2026-01-01,\n"
    first = ingestion.upload("tenant-a", dataset_id="delivery", filename="expected.csv",
                             content=first_bytes)
    duplicate = ingestion.upload("tenant-a", dataset_id="delivery", filename="copy.csv",
                                 content=first_bytes)
    assert duplicate.duplicate and duplicate.document.document_id == first.document.document_id
    replacement = ingestion.upload(
        "tenant-a", dataset_id="delivery", filename="expected.csv",
        content=b"po,supplier,expected,actual\nPO-2,S,2026-01-02,\n")
    assert replacement.document.version == 2
    assert replacement.document.supersedes_document_id == first.document.document_id
    assert not store.get_document("tenant-a", first.document.document_id).active
    ingestion.withdraw("tenant-a", replacement.document.document_id)
    restored = ingestion.upload("tenant-a", dataset_id="delivery", filename="expected.csv",
                                content=first_bytes)
    assert restored.document.version == 3
    assert restored.document.document_id != first.document.document_id
    assert restored.document.supersedes_document_id == replacement.document.document_id


def test_domain_pack_validation_and_authoritative_pomdp_decision(services):
    store, repository, ingestion = services
    ingestion.upload("tenant-a", dataset_id="delivery", filename="expected.csv",
                     content=b"po,supplier,expected,actual\nPO-9,S,2026-01-01,\n")
    ingestion.upload("tenant-a", dataset_id="delivery", filename="actual.csv",
                     content=b"po,supplier,expected,actual\nPO-9,S,,2026-01-04\n")
    candidate = store.list_candidates("tenant-a")[0]
    ingestion.review("tenant-a", candidate.candidate_id, ReviewAction.APPROVE, "reviewer-1")
    packs = DomainPackService(store, repository)
    packs.save("tenant-a", pack())
    scenario = packs.build_scenario("tenant-a", "delivery_delay", "Was PO-9 delayed?")
    response = InvestigationIntelligenceService(
        repository, DeterministicFakeLLM()).investigate_scenario(
            "Was PO-9 delayed?", "tenant-a", scenario)
    assert response.decision == "ANSWER"
    assert response.selected_hypothesis == "delayed"
    assert response.investigation["metadata"]["domain_model_version"] == "reviewed-2026-10"
    assert response.explanation_fallback_used
    assert response.grounding_valid


def test_domain_pack_rejects_unreviewed_probability_tables():
    with pytest.raises(ValidationError):
        pack().model_copy(update={"prior": {"delayed": .9, "on_time": .9}}).model_validate(
            pack().model_copy(update={"prior": {"delayed": .9, "on_time": .9}}).model_dump())


def test_correlated_sources_do_not_satisfy_independence(tmp_path):
    repository = InMemoryGraphRepository()
    store = LocalDocumentStore(tmp_path / "company")
    ingestion = EnterpriseDataService(store, repository)
    ingestion.configure_dataset("tenant-a", dataset())
    ingestion.upload("tenant-a", dataset_id="delivery", filename="expected.csv",
                     content=b"po,supplier,expected,actual\nPO-3,S,2026-01-01,\n",
                     source_group="same-export")
    ingestion.upload("tenant-a", dataset_id="delivery", filename="actual.csv",
                     content=b"po,supplier,expected,actual\nPO-3,S,,2026-01-04\n",
                     source_group="same-export")
    candidate = store.list_candidates("tenant-a")[0]
    ingestion.review("tenant-a", candidate.candidate_id, ReviewAction.APPROVE, "reviewer-1")
    packs = DomainPackService(store, repository)
    packs.save("tenant-a", pack(sources=2))
    scenario = packs.build_scenario("tenant-a", "delivery_delay", "Was PO-3 delayed?")
    response = InvestigationIntelligenceService(
        repository, DeterministicFakeLLM()).investigate_scenario(
            "Was PO-3 delayed?", "tenant-a", scenario)
    assert response.decision == "ABSTAIN"


def test_unsupported_questions_do_not_fall_through_to_company_data(services):
    store, repository, ingestion = services
    with pytest.raises(IngestionError) as lookup_error:
        ingestion.lookup("tenant-a", "Why did every shipment fail?")
    assert lookup_error.value.code == "LOOKUP_UNSUPPORTED"
    packs = DomainPackService(store, repository)
    packs.save("tenant-a", pack())
    with pytest.raises(ValueError, match="purchase order"):
        packs.build_scenario("tenant-a", "delivery_delay", "Was the shipment late?")


def test_on_time_documents_are_contradicting_evidence_for_delay_claim(services):
    store, repository, ingestion = services
    ingestion.upload("tenant-a", dataset_id="delivery", filename="expected.csv",
                     content=b"po,supplier,expected,actual\nPO-ON-TIME,S,2026-01-05,\n")
    ingestion.upload("tenant-a", dataset_id="delivery", filename="actual.csv",
                     content=b"po,supplier,expected,actual\nPO-ON-TIME,S,,2026-01-04\n")
    candidate = store.list_candidates("tenant-a")[0]
    assert candidate.relation_type == "DELIVERED_ON_TIME"
    approved = ingestion.review(
        "tenant-a", candidate.candidate_id, ReviewAction.APPROVE, "reviewer-1")
    supporting, contradicting = repository.get_fact_evidence(
        approved.graph_fact_id, "tenant-a")
    assert supporting == []
    assert len(contradicting) == 2
    packs = DomainPackService(store, repository)
    packs.save("tenant-a", pack())
    scenario = packs.build_scenario(
        "tenant-a", "delivery_delay", "Was PO-ON-TIME delayed?")
    response = InvestigationIntelligenceService(
        repository, DeterministicFakeLLM()).investigate_scenario(
            "Was PO-ON-TIME delayed?", "tenant-a", scenario)
    # Contradiction lowers the delay belief, but cannot become affirmative
    # on-time proof; the evidence gate therefore remains conservative.
    assert response.decision == "ABSTAIN"
    assert response.final_belief["on_time"] > response.final_belief["delayed"]
    assert len(response.contradicting_evidence_ids) == 2


def test_authenticated_api_pilot_and_least_privilege(tmp_path):
    key = secrets.token_urlsafe(36)
    permissions = frozenset({
        Permission.INVESTIGATION_EXECUTE, Permission.EVIDENCE_READ,
        Permission.PROVENANCE_READ, Permission.TRACE_READ,
        Permission.DOCUMENT_UPLOAD, Permission.DOCUMENT_READ,
        Permission.EVIDENCE_REVIEW, Permission.DOMAIN_PACK_MANAGE, Permission.MODEL_APPROVE})
    settings = AppSettings(
        api_key=key, roles=frozenset({"reviewer"}), permissions=permissions,
        document_store_path=str(tmp_path / "api-documents"))
    app = create_app(settings, build_dependencies(settings))
    headers = {"Authorization": f"Bearer {key}"}
    with TestClient(app) as client:
        configured = client.post(
            "/api/v1/company-data/datasets", headers=headers,
            json=dataset().model_dump(mode="json"))
        assert configured.status_code == 200
        for filename, content in (
            ("expected.csv", b"po,supplier,expected,actual\nPO-API,S,2026-01-01,\n"),
            ("actual.csv", b"po,supplier,expected,actual\nPO-API,S,,2026-01-04\n"),
        ):
            response = client.post(
                "/api/v1/company-data/uploads", headers=headers,
                data={"dataset_id": "delivery"},
                files={"file": (filename, content, "text/csv")})
            assert response.status_code == 200
        candidate = client.get(
            "/api/v1/company-data/candidates", headers=headers).json()[0]
        reviewed = client.post(
            f"/api/v1/company-data/candidates/{candidate['candidate_id']}/review",
            headers=headers, json={"action": "APPROVE"})
        assert reviewed.status_code == 200
        lookup = client.post(
            "/api/v1/investigations", headers=headers,
            json={"question": "Was PO-API delayed?", "mode": "document_lookup"})
        assert lookup.status_code == 200
        assert lookup.json()["records"][0]["sources"][0]["source_text"]

        saved_pack = client.post(
            "/api/v1/company-data/domain-packs", headers=headers,
            json=pack().model_copy(update={"status": DomainPackStatus.DRAFT,
                                          "approved_by": None}).model_dump(mode="json"))
        assert saved_pack.status_code == 200
        for action in ("SUBMIT", "APPROVE", "ACTIVATE"):
            lifecycle = client.post(
                f"/api/v1/domains/delivery_delay/{pack().model_version}/lifecycle",
                headers=headers, json={"action": action})
            assert lifecycle.status_code == 200, lifecycle.text
        investigated = client.post(
            "/api/v1/investigations", headers=headers,
            json={"question": "Was PO-API delayed?", "mode": "domain_pomdp",
                  "domain_pack_id": "delivery_delay"})
        assert investigated.status_code == 200
        assert investigated.json()["decision"] == "ANSWER"
        assert investigated.json()["provider"] == "deterministic-domain-pack"

    analyst_key = secrets.token_urlsafe(36)
    analyst_settings = AppSettings(
        api_key=analyst_key, permissions=frozenset({Permission.DOCUMENT_READ}),
        document_store_path=str(tmp_path / "analyst-documents"))
    with TestClient(create_app(analyst_settings, build_dependencies(analyst_settings))) as client:
        denied = client.post(
            "/api/v1/company-data/candidates/not-visible/review",
            headers={"Authorization": f"Bearer {analyst_key}"}, json={"action": "APPROVE"})
        assert denied.status_code == 403
